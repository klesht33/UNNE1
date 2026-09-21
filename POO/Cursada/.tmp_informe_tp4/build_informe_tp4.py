from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

OUT = Path(r"C:\Users\cuent\Documents\UNNE\Clon\POO\Cursada\Informe_TP4_ArandaGabriel.docx")

def set_font(run, size=11, bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold

def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=4):
    p.alignment = align
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = 1.0

def add_text(doc, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=4, size=11):
    p = doc.add_paragraph()
    format_paragraph(p, align, before, after)
    r = p.add_run(text)
    set_font(r, size, bold)
    return p

def add_heading(doc, text, level=1):
    return add_text(doc, text, bold=True, before=8 if level == 1 else 5, after=3)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.5)
section.footer_distance = Inches(0.5)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
normal.font.size = Pt(11)

add_text(doc, "Informe Trabajo Práctico N° 4", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
add_text(doc, "Alumno: Aranda Gabriel", after=2)
add_text(doc, "DNI: 46714612\nProfesores: Mónica Ringa - Daniela Demchum\nGrupo Laboratorio: 1\nTP: N° 4\nFecha de entrega: 21/09/2026", after=8)

add_heading(doc, "1. Introducción y Objetivos")
add_text(doc, "El TP4 aborda el uso de colecciones de Java para administrar grupos de objetos. El trabajo permite reemplazar estructuras creadas manualmente por clases ya disponibles en la biblioteca, concentrando el desarrollo en las reglas del problema: puntos, productos, alumnos, empleados, cuentas y titulares.")
add_text(doc, "El objetivo fue incorporar, eliminar y recuperar elementos de una colección; reconocer las diferencias entre colecciones estáticas y dinámicas, homogéneas y heterogéneas, e indexadas; y elegir la estructura adecuada según la forma de acceso requerida.")

add_heading(doc, "2. Fundamentación Teórica de las Colecciones")
add_heading(doc, "2.1. Reutilización de código", level=2)
add_text(doc, "Las colecciones predefinidas evitan implementar desde cero tareas frecuentes como aumentar la capacidad, recorrer datos, buscar por clave o impedir duplicados. Al instanciar ArrayList, HashMap y HashSet se reutilizan operaciones probadas de la biblioteca de Java, tales como add, remove, get, put, containsKey y size. Esto reduce código repetido y hace más clara la intención de cada clase.")
add_heading(doc, "2.2. Tipos de colecciones", level=2)
add_text(doc, "Un arreglo común es estático porque su capacidad se fija al crearlo. En cambio, ArrayList es dinámico: puede crecer o reducirse durante la ejecución. Además, ArrayList mantiene un orden y permite recuperar un elemento por posición. En el TP, Pedido, Banco y ArrayDePuntos usan este comportamiento. Las colecciones parametrizadas, como ArrayList<Producto> o ArrayList<Empleado>, son homogéneas porque indican qué tipo de objeto almacenan. ArrayDePuntos usa un ArrayList sin genéricos y realiza casting al recuperar; aunque esa forma admite valores de distinto tipo, el programa guarda objetos Punto. HashMap organiza pares clave-valor y HashSet conserva elementos sin repetición.")

add_heading(doc, "3. Análisis de los Ejercicios del TP4")
add_heading(doc, "3.1. Lista de puntos", level=2)
add_text(doc, "ArrayDePuntos crea una colección dinámica de Punto. Desde el programa principal se ingresan seis coordenadas, se agregan objetos mediante agregarPunto y luego se recorren para mostrar sus coordenadas y calcular las distancias entre elementos consecutivos. Al recuperar cada posición se utiliza get y un casting a Punto, evidenciando el acceso indexado de ArrayList.")
add_heading(doc, "3.2. Pedidos y banco con ArrayList", level=2)
add_text(doc, "Pedido mantiene una lista de Producto y ofrece agregarProducto y quitarProducto. TomaPedido agrega unidades desde un catálogo, permite crear productos y recupera el producto elegido con catalogo.get(indice). De modo similar, Banco gestiona listas dinámicas de Empleado y CuentaCorriente: la aplicación agrega, localiza y elimina elementos, y conserva al menos un empleado antes de autorizar una baja.")
add_heading(doc, "3.3. Búsqueda por clave con HashMap", level=2)
add_text(doc, "Curso asocia cada Alumno con su libreta universitaria y Comercio asocia cada Empleado con su CUIL. En ambos casos, put incorpora el objeto, remove lo da de baja, containsKey verifica su existencia y get recupera directamente el valor correspondiente. Este acceso por clave resulta más apropiado que recorrer índices cuando el dato identificador ya está disponible.")
add_heading(doc, "3.4. Titulares sin duplicados", level=2)
add_text(doc, "Banco también utiliza HashSet<Persona> para construir la lista de titulares de sus cuentas. Al agregar cada titular al conjunto, una misma persona no se repite; por ello se obtiene un listado único de clientes aun cuando posean más de una cuenta.")

add_heading(doc, "4. Conclusión")
add_text(doc, "El TP4 demuestra que las colecciones simplifican el manejo de cantidades variables de objetos. ArrayList resuelve listas ordenadas y dinámicas; HashMap facilita búsquedas por clave; y HashSet evita duplicados. La reutilización de estas clases permite que el programa se enfoque en sus objetos y operaciones de negocio, con un código más breve, organizado y fácil de extender.")

doc.core_properties.title = "Informe Trabajo Práctico N° 4"
doc.core_properties.author = "Aranda Gabriel"
doc.save(OUT)
print(OUT)
